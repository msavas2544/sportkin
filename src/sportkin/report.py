"""APA 7 style tables (Word / Markdown) and Bland-Altman figures.

Requires the ``report`` extras: ``pip install "sportkin[report]"``.

APA rules applied: no vertical lines, horizontal rules above and below the
column heads and at the bottom, table number in bold, title in italics,
statistics bounded by 1 (r, ICC, p) printed without a leading zero,
``p < .001`` for very small p values, and a *Note.* line below the table.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np

__all__ = [
    "fmt", "fmt_r", "fmt_p", "fmt_ci",
    "agreement_table", "reliability_table",
    "table_to_markdown", "table_to_docx", "bland_altman_plot",
]

_TXT = {
    "en": {
        "measure": "Measure", "msd": "M ± SD", "bias": "Bias [95% CI]",
        "loa": "95% LoA", "icc": "ICC({m}) [95% CI]", "r": "r", "d": "d",
        "t1": "Trial 1", "t2": "Trial 2", "te": "TE", "cv": "CV%", "mdc": "MDC95",
        "swc": "SWC",
        "note_agree": ("Bias = {a} minus {b}. LoA = limits of agreement (bias ± 1.96 SD of "
                       "the differences). ICC = intraclass correlation coefficient, "
                       "two-way random effects, absolute agreement. d = Cohen's dz (paired samples)."),
        "note_rel": ("TE = typical error; CV% = coefficient of variation; MDC95 = minimal "
                     "detectable change at 95% confidence; SWC = smallest worthwhile change "
                     "(0.2 × between-subject SD)."),
        "mean": "Mean of the two methods", "diff": "Difference ({a} − {b})",
        "table": "Table",
    },
    "tr": {
        "measure": "Ölçüm", "msd": "Ort. ± SS", "bias": "Yanlılık [%95 GA]",
        "loa": "%95 Uyum Sınırları", "icc": "SİKK({m}) [%95 GA]", "r": "r", "d": "d",
        "t1": "1. Ölçüm", "t2": "2. Ölçüm", "te": "TH", "cv": "VK%", "mdc": "MSD95",
        "swc": "EKD",
        "note_agree": ("Yanlılık = {a} eksi {b}. Uyum sınırları = yanlılık ± 1,96 × farkların "
                       "standart sapması. SİKK = sınıf içi korelasyon katsayısı (iki yönlü "
                       "rastgele etkiler, mutlak uyum). d = Cohen dz (eşleştirilmiş örneklem)."),
        "note_rel": ("TH = tipik hata; VK% = varyasyon katsayısı; MSD95 = %95 güvenle "
                     "minimal saptanabilir değişim; EKD = anlamlı en küçük değişim "
                     "(0,2 × denekler arası SS)."),
        "mean": "İki yöntemin ortalaması", "diff": "Fark ({a} − {b})",
        "table": "Tablo",
    },
}


# ----------------------------------------------------------------- formatting
def fmt(x, digits: int = 2, decimal: str = ".") -> str:
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "–"
    s = f"{x:.{digits}f}".replace("-", "−")
    return s.replace(".", decimal)


def fmt_r(x, digits: int = 2, decimal: str = ".") -> str:
    """Correlation-type values without a leading zero (APA, English).

    With a decimal comma (Turkish style) the leading zero is kept: "0,98".
    """
    s = fmt(x, digits, decimal)
    if decimal == "." and abs(x) < 1:
        return s.replace("0" + decimal, decimal, 1)
    return s


def fmt_p(p, decimal: str = ".") -> str:
    if p < 0.001:
        return f"< {decimal}001"
    return "= " + fmt_r(p, 3, decimal)


def fmt_ci(lo, hi, digits: int = 2, decimal: str = ".", r_style: bool = False) -> str:
    f = fmt_r if r_style else fmt
    sep = "; " if decimal == "," else ", "
    return f"[{f(lo, digits, decimal)}{sep}{f(hi, digits, decimal)}]"


# --------------------------------------------------------------------- tables
def agreement_table(results, *, name_a: str = "Video", name_b: str = "Criterion",
                    lang: str = "en", decimal: str = ".", digits: int = 2,
                    title: str = "", number: int = 1) -> dict:
    """Build a method-comparison table.

    results : list of (measure_label, dict returned by ``stats.agreement``)
    """
    T = _TXT[lang]
    m = results[0][1]["icc"]["model"] if results else "2,1"
    head = [T["measure"], f"{name_a} {T['msd']}", f"{name_b} {T['msd']}", T["bias"],
            T["loa"], T["icc"].format(m=m), T["r"], T["d"]]
    rows = []
    for label, res in results:
        ba, ic = res["bland_altman"], res["icc"]
        rows.append([
            label,
            f"{fmt(res['mean_sd_a'][0], digits, decimal)} ± {fmt(res['mean_sd_a'][1], digits, decimal)}",
            f"{fmt(res['mean_sd_b'][0], digits, decimal)} ± {fmt(res['mean_sd_b'][1], digits, decimal)}",
            f"{fmt(ba['bias'], digits, decimal)} {fmt_ci(*ba['bias_ci'], digits, decimal)}",
            fmt_ci(ba["loa_low"], ba["loa_high"], digits, decimal),
            f"{fmt_r(ic['value'], 2, decimal)} {fmt_ci(ic['ci_low'], ic['ci_high'], 2, decimal, True)}",
            fmt_r(res["pearson"]["r"], 2, decimal),
            fmt(res["cohens_d"]["d_z"], 2, decimal),
        ])
    return {"number": number, "title": title, "header": head, "rows": rows,
            "note": T["note_agree"].format(a=name_a, b=name_b), "lang": lang}


def reliability_table(results, *, lang: str = "en", decimal: str = ".", digits: int = 2,
                      title: str = "", number: int = 1) -> dict:
    """Build a test-retest table from ``stats.reliability`` results."""
    T = _TXT[lang]
    m = results[0][1]["icc"]["model"] if results else "2,1"
    head = [T["measure"], f"{T['t1']} {T['msd']}", f"{T['t2']} {T['msd']}",
            T["icc"].format(m=m), T["te"], T["cv"], T["mdc"], T["swc"]]
    rows = []
    for label, res in results:
        ic, te = res["icc"], res["typical_error"]
        rows.append([
            label,
            f"{fmt(res['mean_sd_1'][0], digits, decimal)} ± {fmt(res['mean_sd_1'][1], digits, decimal)}",
            f"{fmt(res['mean_sd_2'][0], digits, decimal)} ± {fmt(res['mean_sd_2'][1], digits, decimal)}",
            f"{fmt_r(ic['value'], 2, decimal)} {fmt_ci(ic['ci_low'], ic['ci_high'], 2, decimal, True)}",
            fmt(te["typical_error"], digits, decimal),
            fmt(te["cv_percent"], 1, decimal),
            fmt(te["mdc95"], digits, decimal),
            fmt(te["swc"], digits, decimal),
        ])
    return {"number": number, "title": title, "header": head, "rows": rows,
            "note": T["note_rel"], "lang": lang}


def table_to_markdown(table: dict) -> str:
    T = _TXT[table["lang"]]
    out = [f"**{T['table']} {table['number']}**", ""]
    if table["title"]:
        out += [f"*{table['title']}*", ""]
    out.append("| " + " | ".join(table["header"]) + " |")
    out.append("|" + "|".join("---" for _ in table["header"]) + "|")
    out += ["| " + " | ".join(r) + " |" for r in table["rows"]]
    note = "Not." if table["lang"] == "tr" else "Note."
    out += ["", f"*{note}* {table['note']}"]
    return "\n".join(out)


def _set_borders(cell, top=None, bottom=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, on in (("top", top), ("bottom", bottom), ("left", False), ("right", False)):
        if on is None:
            continue
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if on else "nil")
        if on:
            el.set(qn("w:sz"), "8")
            el.set(qn("w:color"), "000000")
        borders.append(el)


def table_to_docx(tables, path, font: str = "Times New Roman", size_pt: int = 10) -> Path:
    """Write one or more APA tables to a Word document."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    if isinstance(tables, dict):
        tables = [tables]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size_pt)
    if max(len(tb["header"]) for tb in tables) > 6:
        from docx.enum.section import WD_ORIENT

        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for k, tb in enumerate(tables):
        T = _TXT[tb["lang"]]
        if k:
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run(f"{T['table']} {tb['number']}").bold = True
        p.paragraph_format.space_after = Pt(0)
        if tb["title"]:
            doc.add_paragraph().add_run(tb["title"]).italic = True
        n_cols = len(tb["header"])
        t = doc.add_table(rows=1 + len(tb["rows"]), cols=n_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        all_rows = [tb["header"]] + tb["rows"]
        for i, row in enumerate(all_rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(val))
                run.font.size = Pt(size_pt)
                para = cell.paragraphs[0]
                para.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0
                                  else WD_ALIGN_PARAGRAPH.CENTER)
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                last = i == len(all_rows) - 1
                _set_borders(cell, top=(i <= 1), bottom=(i == 0 or last))
        note = doc.add_paragraph()
        note.add_run("Not." if tb["lang"] == "tr" else "Note.").italic = True
        note.add_run(" " + tb["note"])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# --------------------------------------------------------------------- figure
def bland_altman_plot(ba, path=None, *, name_a: str = "Video", name_b: str = "Criterion",
                      unit: str = "cm", lang: str = "en", dpi: int = 300, show_ci: bool = True):
    """Grey-scale Bland-Altman plot suitable for journals.

    ``ba`` is a :class:`sportkin.stats.BlandAltmanResult`.
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    T = _TXT[lang]
    dec = "," if lang == "tr" else "."
    fig, ax = plt.subplots(figsize=(5.5, 4.0))
    ax.scatter(ba.means, ba.diffs, s=22, facecolor="white", edgecolor="black", linewidth=0.9, zorder=3)
    x0, x1 = np.min(ba.means), np.max(ba.means)
    pad = (x1 - x0) * 0.08 or 1
    xs = (x0 - pad, x1 + pad)
    lines = [(ba.bias, "-", ba.bias_ci), (ba.loa_low, "--", ba.loa_low_ci),
             (ba.loa_high, "--", ba.loa_high_ci)]
    for y, style, ci in lines:
        ax.plot(xs, (y, y), style, color="black", linewidth=1)
        if show_ci:
            ax.axhspan(ci[0], ci[1], color="0.85", zorder=0)
        ax.annotate(f"{y:.2f}".replace(".", dec).replace("-", "−"), xy=(xs[1], y), xytext=(4, 0),
                    textcoords="offset points", va="center", fontsize=8)
    ax.axhline(0, color="0.6", linewidth=0.6, zorder=1)
    if dec == ",":
        from matplotlib.ticker import FuncFormatter

        tick = FuncFormatter(lambda v, _: f"{v:g}".replace(".", ",").replace("-", "−"))
        ax.xaxis.set_major_formatter(tick)
        ax.yaxis.set_major_formatter(tick)
    ax.set_xlim(*xs)
    u = f" ({unit})" if unit else ""
    ax.set_xlabel(T["mean"] + u)
    ax.set_ylabel(T["diff"].format(a=name_a, b=name_b) + u)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    if path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(path, dpi=dpi)
        plt.close(fig)
        return path
    return fig
